from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .models import PageData


class ExcelExporter:
    @staticmethod
    def export(pages: dict[str, PageData], path: str = "results.xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Crawl Results"

        headers = ["URL", "Status Code", "Inbound Links", "Outbound Links"]
        header_font = Font(bold=True, size=11)
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        row = 2
        for url in sorted(pages.keys()):
            page = pages[url]
            ws.cell(row=row, column=1, value=url)
            ws.cell(row=row, column=2, value=page.status_code)
            ws.cell(
                row=row, column=3,
                value=ExcelExporter._format_links(page.inbound, pages)
            )
            ws.cell(
                row=row, column=4,
                value=ExcelExporter._format_links(page.outbound, pages)
            )
            row += 1

        ws.column_dimensions["A"].width = 60
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 80
        ws.column_dimensions["D"].width = 80

        wb.save(path)
        print(f"Saved: {path}")

    @staticmethod
    def _format_links(links: set[str],
                      pages: dict[str, PageData]) -> str:
        if not links:
            return ""
        formatted = []
        for link in sorted(links):
            status = (
                pages[link].status_code if link in pages else "?"
            )
            formatted.append(f"{link} ({status})")
        return ", ".join(formatted)
